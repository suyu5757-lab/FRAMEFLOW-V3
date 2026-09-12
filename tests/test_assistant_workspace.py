from __future__ import annotations

import asyncio
import re
import shutil
import tempfile
import time
import uuid
import unittest
from pathlib import Path
from unittest import mock

from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import server
from frameflow.assistant_attachments import extract_document


def _project(project_id: str) -> dict:
    return {
        "id": project_id,
        "name": "Agent 工作台测试",
        "ratio": "16:9",
        "duration": 8,
        "generator": "V3 local",
        "brief": "测试项目",
        "stage": 0,
        "sortOrder": 0,
        "script": "",
        "assets": [{
            "id": "AST_AGENT",
            "name": "测试角色",
            "assetClass": "character",
            "type": "角色",
            "grade": "B",
            "status": "missing",
        }],
        "shots": [],
        "audio": {},
        "assetRegulator": {},
        "generations": [],
        "seedancePackages": [],
        "providerOverrides": {},
        "undoStack": [],
        "scriptVersions": [],
        "storyboardVersions": [],
        "storyWorkflowRuns": [],
    }


class FakeAgentAdapter:
    def __init__(self, response: dict | None = None, capabilities: set[str] | None = None) -> None:
        self.calls: list[tuple[str, dict, str]] = []
        self.response = response
        self.capabilities = capabilities or {"orchestrator"}
        self.error: Exception | None = None
        self.delay = 0.0

    def supports(self, capability: str) -> bool:
        return capability in self.capabilities

    def validate_request(self, capability: str, request: dict) -> list[str]:
        return []

    async def submit(self, capability: str, request: dict, credential: str) -> dict:
        self.calls.append((capability, request, credential))
        if self.delay:
            await asyncio.sleep(self.delay)
        if self.error is not None:
            raise self.error
        return self.response or {
            "structured": {
                "reply": "已读取资料并生成候选。",
                "patch": {
                    "workspace_operations": [{
                        "id": "OP_AGENT_PROMPT",
                        "workspace": "assets",
                        "action": "create_prompt_candidate",
                        "target_id": "AST_AGENT",
                        "title": "角色 Prompt 候选",
                        "summary": "保留稳定角色 ID，补充可见身份锚点。",
                        "after": {
                            "prompt": "一名角色的正面中景肖像，保持脸部身份锚点，冷色侧光在服装材质上留下清晰高光。",
                            "promptPack": {"promptIntent": "建立角色身份参考"},
                        },
                    }],
                },
                "actions": ["candidate_draft"],
                "next_skill": None,
                "requires_confirmation": False,
            },
        }


class AssistantWorkspaceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_root = Path(tempfile.mkdtemp(prefix="frameflow-assistant-test-"))
        self.db_path = self.temp_root / ("db-" + uuid.uuid4().hex + ".sqlite")
        self.project_id = "PRJ_AGENT_" + uuid.uuid4().hex[:8].upper()
        self.adapter = FakeAgentAdapter()
        self.db_patch = mock.patch.object(server, "DB_PATH", self.db_path)
        self.secret_patch = mock.patch.object(server, "get_secret", return_value="test-secret")
        self.adapter_patch = mock.patch.object(server, "adapter_for_profile", return_value=self.adapter)
        self.profile_secret_patch = mock.patch.object(server, "get_profile_secret", return_value="test-secret")
        self.db_patch.start()
        self.secret_patch.start()
        self.adapter_patch.start()
        self.profile_secret_patch.start()
        self.client_context = TestClient(server.app)
        self.client = self.client_context.__enter__()
        created = self.client.put("/api/v2/projects/" + self.project_id, json={"document": _project(self.project_id)})
        assert created.status_code == 200, created.text

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        self.profile_secret_patch.stop()
        self.adapter_patch.stop()
        self.secret_patch.stop()
        self.db_patch.stop()
        runtime_root = Path(tempfile.gettempdir()) / ("frameflow-runtime-" + self.db_path.stem)
        if runtime_root.is_dir():
            shutil.rmtree(runtime_root)
        if self.temp_root.is_dir():
            shutil.rmtree(self.temp_root)

    def _upload(self, filename: str, content: bytes, mime_type: str, conversation_id: str | None = None) -> dict:
        query = f"?conversation_id={conversation_id}" if conversation_id else ""
        response = self.client.post(
            f"/api/v2/projects/{self.project_id}/assistant/attachments{query}",
            files={"file": (filename, content, mime_type)},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["attachment"]

    def _start_run(
        self,
        message: str = "检查当前项目并生成可审阅候选",
        attachment_ids: list[str] | None = None,
        conversation_id: str | None = None,
        client_message_id: str | None = None,
        **overrides: object,
    ) -> str:
        body: dict[str, object] = {
            "project_id": self.project_id,
            "conversation_id": conversation_id,
            "message": message,
            "attachment_ids": attachment_ids or [],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {"confirmation_required": True},
            "client_message_id": client_message_id or "client-" + uuid.uuid4().hex,
        }
        body.update(overrides)
        response = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json=body)
        self.assertEqual(response.status_code, 200, response.text)
        match = re.search(r'"run_id":\s*"([^"]+)"', response.text)
        self.assertIsNotNone(match, response.text)
        return str(match.group(1))

    def _wait_for_terminal(self, run_id: str, timeout: float = 4.0) -> dict:
        deadline = time.monotonic() + timeout
        run: dict = self.client.get(f"/api/v2/assistant/runs/{run_id}").json()
        while run.get("status") not in {"succeeded", "failed", "canceled", "stale_contract"} and time.monotonic() < deadline:
            time.sleep(0.05)
            run = self.client.get(f"/api/v2/assistant/runs/{run_id}").json()
        self.assertIn(run.get("status"), {"succeeded", "failed", "canceled", "stale_contract"}, run)
        return run

    def test_explicit_resource_dir_errors_instead_of_falling_back(self) -> None:
        missing = self.temp_root / "resource-does-not-exist"
        with mock.patch.object(server, "configured_resource_dir", str(missing)):
            with self.assertRaises(RuntimeError):
                server.validate_configured_resource_dir()

    def test_archiving_only_changes_conversation_state_and_keeps_attachment(self) -> None:
        created = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/conversations", json={"title": "可归档会话"})
        assert created.status_code == 200, created.text
        conversation_id = created.json()["conversation"]["id"]
        uploaded = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/attachments?conversation_id=" + conversation_id,
            files={"file": ("archive-note.txt", b"keep this material", "text/plain")},
        )
        assert uploaded.status_code == 200, uploaded.text
        attachment_url = uploaded.json()["attachment"]["url"]
        archived = self.client.post("/api/v2/assistant/conversations/" + conversation_id + "/archive", json={})
        assert archived.status_code == 200, archived.text
        assert archived.json()["conversation"]["status"] == "archived"
        assert self.client.get(attachment_url).content == b"keep this material"
        restored = self.client.post("/api/v2/assistant/conversations/" + conversation_id + "/restore", json={})
        assert restored.status_code == 200, restored.text
        assert restored.json()["conversation"]["status"] == "active"

    def test_contract_bundle_and_local_attachment_are_project_scoped(self) -> None:
        contracts = self.client.get("/api/v2/contracts")
        assert contracts.status_code == 200, contracts.text
        payload = contracts.json()
        assert payload["bundle_hash"]
        assert payload["prompt_contract"]["version"]
        assert payload["story_contract"]["required_shot_fields"]
        assert payload["audio_contract"]["required_fields"]
        settings = self.client.get("/api/v2/settings")
        assert settings.status_code == 200, settings.text
        assert settings.json()["feature_flags"]["assistant_workspace_v2"] is True

        uploaded = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/attachments",
            files={"file": ("notes.txt", b"FRAMEFLOW local notes", "text/plain")},
        )
        assert uploaded.status_code == 200, uploaded.text
        attachment = uploaded.json()["attachment"]
        assert attachment["delivery_mode"] == "pending"
        assert attachment["url"].startswith("/api/v2/assistant/attachments/")
        assert "storage_path" not in attachment
        served = self.client.get(attachment["url"])
        assert served.status_code == 200
        assert served.content == b"FRAMEFLOW local notes"

    def test_contract_scope_workflow_and_legacy_routes_are_explicit(self) -> None:
        bundle = self.client.get("/api/v2/contracts")
        self.assertEqual(bundle.status_code, 200, bundle.text)
        bundle_payload = bundle.json()
        for scope in ("prompt", "story", "audio", "workflow"):
            response = self.client.get(f"/api/v2/contracts/{scope}")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["bundle_hash"], bundle_payload["bundle_hash"])
        workflows = self.client.get("/api/v2/workflows")
        self.assertEqual(workflows.status_code, 200, workflows.text)
        self.assertEqual(workflows.json()["contract_hash"], bundle_payload["bundle_hash"])
        self.assertTrue(workflows.json()["workflows"])
        unknown = self.client.get("/api/v2/contracts/does-not-exist")
        self.assertEqual(unknown.status_code, 404, unknown.text)
        legacy = self.client.post("/api/assistant/stream", json={})
        self.assertEqual(legacy.status_code, 410, legacy.text)

    def test_provider_instructions_use_live_contract_and_exclude_local_paths(self) -> None:
        run_id = self._start_run("检查实时规范和隐私边界", context={"file_path": "/private/should-never-be-sent"})
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        request = self.adapter.calls[-1][1]
        contracts = self.client.get("/api/v2/contracts").json()
        self.assertIn(contracts["bundle_hash"], request["instructions"])
        self.assertIn("永不执行", request["instructions"])
        self.assertNotIn("/private/should-never-be-sent", request["input_text"])
        self.assertNotIn(str(server.DATA_DIR), request["input_text"])
        self.assertNotIn("test-secret", request["input_text"])

    def test_large_project_and_attachment_context_are_bounded_without_rewriting_project(self) -> None:
        current = self.client.get(f"/api/v2/projects/{self.project_id}").json()["document"]
        current["assets"][0].update({
            "prompt": "CURRENT-PROMPT " + "p" * 16000,
            "promptPack": {"promptIntent": "保留当前资产稳定身份", "identityAnchor": "CURRENT-ANCHOR " + "a" * 6000},
        })
        current["assetPromptRuns"] = [{"id": "HISTORY_RUN", "status": "prompt_drafts_ready", "prompt": "HISTORY-PROMPT " + "h" * 260000}]
        current["storyboardVersions"] = [{"id": "HISTORY_STORY", "status": "superseded", "package": {"shots": [{"id": "S" + str(index), "action": "history " + "s" * 3000} for index in range(20)]}}]
        database = server.app.state.db
        with database.connect() as connection:
            row = connection.execute("SELECT revision FROM projects WHERE id=?", (self.project_id,)).fetchone()
            connection.execute("UPDATE projects SET document_json=?,revision=?,updated_at=? WHERE id=?", (database.encode(current), int(row["revision"]) + 1, server.utcnow(), self.project_id))

        first = self._upload("large-one.txt", ("first document " + "一" * 40000).encode("utf-8"), "text/plain")
        second = self._upload("large-two.txt", ("second document " + "二" * 40000).encode("utf-8"), "text/plain")
        run_id = self._start_run("请按当前工作台状态做一次简单测试", [first["id"], second["id"]])
        waiting = self.client.get(f"/api/v2/assistant/runs/{run_id}").json()
        self.assertEqual(waiting["status"], "awaiting_external_confirmation", waiting)
        resumed = self.client.post(f"/api/v2/assistant/runs/{run_id}/external-confirmation", json={"decision": "approve"})
        self.assertEqual(resumed.status_code, 200, resumed.text)
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        request = self.adapter.calls[-1][1]
        self.assertLessEqual(len(request["input_text"]), server.ASSISTANT_INPUT_TEXT_SAFE_LIMIT)
        self.assertIn("bounded-current-state-v1", request["input_text"])
        self.assertIn("附件引用", request["input_text"])
        self.assertIn("已省略", request["input_text"])
        self.assertNotIn("HISTORY-PROMPT " + "h" * 1000, request["input_text"])
        progress = [event for event in self.client.get(f"/api/v2/assistant/runs/{run_id}/events").json()["events"] if event["item_id"] == "context_loading" and event["event_type"] == "item_progress"]
        self.assertTrue(progress)
        self.assertTrue(progress[-1]["data"]["project_context_compacted"])

        unchanged = self.client.get(f"/api/v2/projects/{self.project_id}").json()["document"]
        self.assertEqual(len(unchanged["assetPromptRuns"][0]["prompt"]), 260015)
        self.assertEqual(len(unchanged["storyboardVersions"][0]["package"]["shots"]), 20)

    def test_conversations_are_project_isolated_and_archived_conversations_block_send(self) -> None:
        other_project_id = "PRJ_OTHER_" + uuid.uuid4().hex[:8].upper()
        created = self.client.put(f"/api/v2/projects/{other_project_id}", json={"document": _project(other_project_id)})
        self.assertEqual(created.status_code, 200, created.text)
        own = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/conversations", json={"title": "当前项目会话"})
        other = self.client.post(f"/api/v2/projects/{other_project_id}/assistant/conversations", json={"title": "另一个项目会话"})
        self.assertEqual(own.status_code, 200, own.text)
        self.assertEqual(other.status_code, 200, other.text)
        own_id = own.json()["conversation"]["id"]
        other_id = other.json()["conversation"]["id"]
        own_list = self.client.get(f"/api/v2/projects/{self.project_id}/assistant/conversations")
        self.assertEqual(own_list.status_code, 200, own_list.text)
        self.assertEqual([item["id"] for item in own_list.json()["conversations"]], [own_id])
        wrong_upload = self.client.post(
            f"/api/v2/projects/{self.project_id}/assistant/attachments?conversation_id={other_id}",
            files={"file": ("wrong-project.txt", b"must reject", "text/plain")},
        )
        self.assertEqual(wrong_upload.status_code, 404, wrong_upload.text)
        wrong_stream = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": self.project_id,
            "conversation_id": other_id,
            "message": "不能读取另一个项目会话",
            "attachment_ids": [],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "isolation-" + uuid.uuid4().hex,
        })
        self.assertEqual(wrong_stream.status_code, 404, wrong_stream.text)
        archived = self.client.post(f"/api/v2/assistant/conversations/{own_id}/archive", json={})
        self.assertEqual(archived.status_code, 200, archived.text)
        blocked = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": self.project_id,
            "conversation_id": own_id,
            "message": "归档后不能发送",
            "attachment_ids": [],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "archived-" + uuid.uuid4().hex,
        })
        self.assertEqual(blocked.status_code, 409, blocked.text)
        restored = self.client.post(f"/api/v2/assistant/conversations/{own_id}/restore", json={})
        self.assertEqual(restored.status_code, 200, restored.text)
        self.assertEqual(restored.json()["conversation"]["status"], "active")

    def test_conversation_title_can_be_renamed_and_exposes_last_message_preview(self) -> None:
        created = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/conversations", json={"title": "原始标题"})
        self.assertEqual(created.status_code, 200, created.text)
        conversation_id = created.json()["conversation"]["id"]
        renamed = self.client.patch(f"/api/v2/assistant/conversations/{conversation_id}", json={"title": "人工整理后的标题"})
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual(renamed.json()["conversation"]["title"], "人工整理后的标题")
        empty_title = self.client.patch(f"/api/v2/assistant/conversations/{conversation_id}", json={"title": "   "})
        self.assertEqual(empty_title.status_code, 422, empty_title.text)
        run_id = self._start_run("产生一条可以搜索的最近消息", conversation_id=conversation_id)
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        listed = self.client.get(f"/api/v2/projects/{self.project_id}/assistant/conversations").json()["conversations"]
        current = next(item for item in listed if item["id"] == conversation_id)
        self.assertEqual(current["title"], "人工整理后的标题")
        self.assertTrue(current["last_message"])

    def test_attachment_limits_duplicate_references_and_path_safety(self) -> None:
        empty = self.client.post(
            f"/api/v2/projects/{self.project_id}/assistant/attachments",
            files={"file": ("empty.txt", b"", "text/plain")},
        )
        self.assertEqual(empty.status_code, 422, empty.text)
        traversal = self.client.post(
            f"/api/v2/projects/{self.project_id}/assistant/attachments",
            files={"file": ("../../escape.txt", b"must reject", "text/plain")},
        )
        self.assertEqual(traversal.status_code, 422, traversal.text)
        safe = self._upload("escape.txt", b"safe content", "text/plain")
        self.assertEqual(safe["safe_name"], "escape.txt")
        attachment_dir = server.DATA_DIR / "projects" / self.project_id / "assistant" / "attachments"
        self.assertEqual(Path(safe["url"]).name, safe["id"])
        self.assertTrue((attachment_dir / f"{safe['id']}.txt").is_file())
        before_oversize = {item.name for item in attachment_dir.glob("*")}
        with mock.patch.object(server, "MAX_ATTACHMENT_BYTES", 4):
            oversize = self.client.post(
                f"/api/v2/projects/{self.project_id}/assistant/attachments",
                files={"file": ("too-large.txt", b"12345", "text/plain")},
            )
        self.assertEqual(oversize.status_code, 413, oversize.text)
        self.assertEqual({item.name for item in attachment_dir.glob("*")}, before_oversize)
        attachments = [safe]
        for index in range(8):
            attachments.append(self._upload(f"notes-{index}.txt", f"note {index}".encode(), "text/plain"))
        ids = [item["id"] for item in attachments]
        too_many = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": self.project_id,
            "message": "不能超过单条消息附件数量上限",
            "attachment_ids": ids,
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "too-many-" + uuid.uuid4().hex,
        })
        self.assertEqual(too_many.status_code, 422, too_many.text)
        duplicate = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": self.project_id,
            "message": "不能重复引用附件",
            "attachment_ids": [safe["id"], safe["id"]],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "duplicate-" + uuid.uuid4().hex,
        })
        self.assertEqual(duplicate.status_code, 422, duplicate.text)
        missing = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": self.project_id,
            "message": "不能引用不存在的附件",
            "attachment_ids": ["ATT_DOES_NOT_EXIST"],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "missing-" + uuid.uuid4().hex,
        })
        self.assertEqual(missing.status_code, 404, missing.text)
        with mock.patch.object(server, "MAX_MESSAGE_BYTES", 4):
            total_oversize = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
                "project_id": self.project_id,
                "message": "不能超过单条消息总大小",
                "attachment_ids": [safe["id"]],
                "selected_node_ids": [],
                "skill_id": "video-script-storyboard",
                "context": {},
                "cost_boundary": {},
                "client_message_id": "total-" + uuid.uuid4().hex,
            })
        self.assertEqual(total_oversize.status_code, 413, total_oversize.text)

    def test_audio_video_and_subtitle_attachments_remain_project_references(self) -> None:
        audio = self._upload("voice.mp3", b"audio bytes", "audio/mpeg")
        video = self._upload("reference.mp4", b"video bytes", "video/mp4")
        subtitle = self._upload("dialogue.srt", "1\n00:00:00,000 --> 00:00:01,000\n你好\n".encode(), "text/plain")
        run_id = self._start_run("只检查本地资料引用", [audio["id"], video["id"], subtitle["id"]])
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        self.assertEqual({item["delivery_mode"] for item in run["attachments"]}, {"project_reference"})
        self.assertEqual([item[0] for item in self.adapter.calls], ["orchestrator"])
        prompt_text = self.adapter.calls[0][1]["input_text"]
        self.assertIn("仅作为本地项目资料引用", prompt_text)
        self.assertNotIn(str(server.DATA_DIR), prompt_text)
        events = self.client.get(f"/api/v2/assistant/runs/{run_id}/events").json()["events"]
        self.assertFalse(any(item["event_type"] == "approval_request" for item in events))

    def test_provider_without_vision_keeps_image_local_and_marks_analysis_unavailable(self) -> None:
        image = self._upload("no-vision.png", b"fake image", "image/png")
        run_id = self._start_run("当前 Provider 不支持 vision", [image["id"]])
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        self.assertEqual([item[0] for item in self.adapter.calls], ["orchestrator"])
        stored = next(item for item in run["attachments"] if item["id"] == image["id"])
        self.assertEqual(stored["delivery_mode"], "project_reference")
        self.assertEqual(stored["analysis_status"], "unavailable")
        events = self.client.get(f"/api/v2/assistant/runs/{run_id}/events").json()["events"]
        vision_events = [item for item in events if item["item_id"] == "vision_analysis"]
        self.assertTrue(vision_events, events)
        self.assertEqual(vision_events[-1]["status"], "skipped")
        self.assertEqual(run["result"].get("vision_analyzed_ids"), [])

    def test_external_confirmation_reject_reuse_and_reset_follow_provider_scope(self) -> None:
        first_attachment = self._upload("first.txt", b"first", "text/plain")
        first_run_id = self._start_run("第一次需要确认", [first_attachment["id"]])
        first_run = self.client.get(f"/api/v2/assistant/runs/{first_run_id}").json()
        self.assertEqual(first_run["status"], "awaiting_external_confirmation")
        conversation_id = first_run["conversation_id"]
        wrong_provider = self.client.post(f"/api/v2/assistant/runs/{first_run_id}/external-confirmation", json={
            "decision": "approve",
            "provider_profile_id": "openai-default",
            "detail": {"approved_by": "test"},
        })
        self.assertEqual(wrong_provider.status_code, 409, wrong_provider.text)
        rejected = self.client.post(f"/api/v2/assistant/runs/{first_run_id}/external-confirmation", json={
            "decision": "reject",
            "provider_profile_id": "opencode-default",
            "detail": {"approved_by": "test"},
        })
        self.assertEqual(rejected.status_code, 200, rejected.text)
        self.assertEqual(rejected.json()["status"], "canceled")
        self.assertEqual(self.adapter.calls, [])
        self.assertEqual(self.client.get(first_attachment["url"]).content, b"first")

        second_attachment = self._upload("second.txt", b"second", "text/plain", conversation_id)
        second_run_id = self._start_run("同一会话第二次确认", [second_attachment["id"]], conversation_id)
        self.assertEqual(self.client.get(f"/api/v2/assistant/runs/{second_run_id}").json()["status"], "awaiting_external_confirmation")
        approved = self.client.post(f"/api/v2/assistant/runs/{second_run_id}/external-confirmation", json={
            "decision": "approve",
            "provider_profile_id": "opencode-default",
            "detail": {"approved_by": "test"},
        })
        self.assertEqual(approved.status_code, 200, approved.text)
        second_run = self._wait_for_terminal(second_run_id)
        self.assertEqual(second_run["status"], "succeeded", second_run)
        self.assertEqual([item[0] for item in self.adapter.calls], ["orchestrator"])

        third_attachment = self._upload("third.txt", b"third", "text/plain", conversation_id)
        third_run_id = self._start_run("同 Provider 可复用确认", [third_attachment["id"]], conversation_id)
        third_run = self._wait_for_terminal(third_run_id)
        self.assertEqual(third_run["status"], "succeeded", third_run)
        self.assertEqual(len(self.adapter.calls), 2)
        consented = self.client.get(f"/api/v2/projects/{self.project_id}/assistant/conversations").json()["conversations"]
        self.assertIn("opencode-default", next(item for item in consented if item["id"] == conversation_id)["external_consent"])
        switched_attachment = self._upload("switched-provider.txt", b"switch", "text/plain", conversation_id)
        switched_run_id = self._start_run(
            "切换 Provider 后必须重新确认",
            [switched_attachment["id"]],
            conversation_id,
            provider_profile_id="openai-default",
            model="gpt-5.6-terra",
        )
        self.assertEqual(self.client.get(f"/api/v2/assistant/runs/{switched_run_id}").json()["status"], "awaiting_external_confirmation")
        switched_rejected = self.client.post(f"/api/v2/assistant/runs/{switched_run_id}/external-confirmation", json={
            "decision": "reject",
            "provider_profile_id": "openai-default",
            "detail": {"approved_by": "test"},
        })
        self.assertEqual(switched_rejected.status_code, 200, switched_rejected.text)
        self.assertEqual(len(self.adapter.calls), 2)
        reset = self.client.post(f"/api/v2/assistant/conversations/{conversation_id}/external-consent/reset", json={})
        self.assertEqual(reset.status_code, 200, reset.text)
        self.assertEqual(reset.json()["conversation"]["external_consent"], [])
        fourth_attachment = self._upload("fourth.txt", b"fourth", "text/plain", conversation_id)
        fourth_run_id = self._start_run("重置后必须再次确认", [fourth_attachment["id"]], conversation_id)
        self.assertEqual(self.client.get(f"/api/v2/assistant/runs/{fourth_run_id}").json()["status"], "awaiting_external_confirmation")
        self.assertEqual(len(self.adapter.calls), 2)

    def test_run_events_messages_plans_and_sequence_replay_survive_refresh(self) -> None:
        run_id = self._start_run("验证刷新后的持久化状态")
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        all_events = self.client.get(f"/api/v2/assistant/runs/{run_id}/events").json()["events"]
        sequences = [item["sequence"] for item in all_events]
        self.assertEqual(sequences, list(range(1, len(sequences) + 1)))
        midpoint = sequences[len(sequences) // 2]
        replay = self.client.get(f"/api/v2/assistant/runs/{run_id}/events?after_sequence={midpoint}")
        self.assertEqual(replay.status_code, 200, replay.text)
        self.assertTrue(all(item["sequence"] > midpoint for item in replay.json()["events"]))
        self.assertTrue(any(item["event_type"] == "run_completed" for item in all_events))
        messages = self.client.get(f"/api/v2/assistant/conversations/{run['conversation_id']}/messages")
        self.assertEqual(messages.status_code, 200, messages.text)
        message_payload = messages.json()["messages"]
        self.assertEqual([item["role"] for item in message_payload], ["user", "assistant"])
        self.assertEqual(message_payload[-1]["metadata"]["run_id"], run_id)
        runs = self.client.get(f"/api/v2/projects/{self.project_id}/assistant/runs?conversation_id={run['conversation_id']}")
        self.assertEqual(runs.status_code, 200, runs.text)
        self.assertEqual(runs.json()["runs"][0]["id"], run_id)
        plan_events = self.client.get(f"/api/v2/agent/plans/{run['result']['plan_id']}/events")
        self.assertEqual(plan_events.status_code, 200, plan_events.text)
        self.assertTrue(any(item["event"] == "created" for item in plan_events.json()["events"]))

    def test_reject_plan_preserves_project_and_history(self) -> None:
        run_id = self._start_run("生成一个待拒绝的候选")
        run = self._wait_for_terminal(run_id)
        before = self.client.get(f"/api/v2/projects/{self.project_id}").json()
        rejected = self.client.post(f"/api/v2/assistant/runs/{run_id}/reject", json={"detail": {"reason": "不符合当前方向"}})
        self.assertEqual(rejected.status_code, 200, rejected.text)
        plan = self.client.get(f"/api/v2/agent/plans/{run['result']['plan_id']}").json()
        self.assertEqual(plan["status"], "rejected")
        self.assertEqual(plan["decision"]["detail"]["reason"], "不符合当前方向")
        after = self.client.get(f"/api/v2/projects/{self.project_id}").json()
        self.assertEqual(after["revision"], before["revision"])
        self.assertEqual(after["document"], before["document"])
        messages = self.client.get(f"/api/v2/assistant/conversations/{run['conversation_id']}/messages").json()["messages"]
        self.assertEqual(len(messages), 2)

    def test_apply_requires_selection_and_rejects_revision_conflict_before_write(self) -> None:
        run_id = self._start_run("验证应用前的乐观并发门禁")
        run = self._wait_for_terminal(run_id)
        plan_id = run["result"]["plan_id"]
        operation_id = run["result"]["patch"]["workspace_operations"][0]["id"]
        contracts = self.client.get("/api/v2/contracts").json()
        empty_selection = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": plan_id,
            "selected_operation_ids": [],
            "expected_project_revision": run["base_project_revision"],
            "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"],
            "expected_contract_bundle_hash": contracts["bundle_hash"],
        })
        self.assertEqual(empty_selection.status_code, 422, empty_selection.text)
        before = self.client.get(f"/api/v2/projects/{self.project_id}").json()
        conflict = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": plan_id,
            "selected_operation_ids": [operation_id],
            "expected_project_revision": run["base_project_revision"] + 1,
            "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"],
            "expected_contract_bundle_hash": contracts["bundle_hash"],
        })
        self.assertEqual(conflict.status_code, 409, conflict.text)
        after = self.client.get(f"/api/v2/projects/{self.project_id}").json()
        self.assertEqual(after, before)
        self.assertEqual(self.client.get(f"/api/v2/agent/plans/{plan_id}").json()["status"], "awaiting_review")
        applied = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": plan_id,
            "selected_operation_ids": [operation_id],
            "expected_project_revision": run["base_project_revision"],
            "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"],
            "expected_contract_bundle_hash": contracts["bundle_hash"],
        })
        self.assertEqual(applied.status_code, 200, applied.text)
        self.assertEqual(applied.json()["plan"]["status"], "applied")

    def test_blocked_workflow_operation_cannot_be_applied_by_bypassing_frontend(self) -> None:
        self.adapter.response = {"structured": {
            "reply": "建议移除一个节点。",
            "patch": {"remove_node_ids": ["generate"], "actions": ["node_orchestration"]},
        }}
        run_id = self._start_run("尝试验证受保护工作流操作")
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        operation = next(item for item in run["result"]["patch"]["workspace_operations"] if item["workspace"] == "workflow")
        self.assertEqual(operation["risk"], "blocked")
        before_graph = self.client.get(f"/api/v2/projects/{self.project_id}/graph").json()
        response = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"],
            "selected_operation_ids": [operation["id"]],
            "expected_project_revision": run["base_project_revision"],
            "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"],
            "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
        })
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(self.client.get(f"/api/v2/projects/{self.project_id}/graph").json(), before_graph)
        self.assertEqual(self.client.get(f"/api/v2/agent/plans/{run['result']['plan_id']}").json()["status"], "awaiting_review")

    def test_apply_transaction_rolls_back_candidate_rows_when_a_later_operation_fails(self) -> None:
        self.adapter.response = {"structured": {
            "reply": "包含一项有效元数据候选和一项无效目标候选。",
            "patch": {"workspace_operations": [
                {"id": "OP_METADATA_OK", "workspace": "assets", "action": "update_metadata", "target_id": "AST_AGENT", "title": "有效元数据", "summary": "保留身份锚点。", "content": {"mustPreserve": ["面部身份"]}},
                {"id": "OP_METADATA_MISSING", "workspace": "assets", "action": "update_metadata", "target_id": "AST_MISSING", "title": "无效目标", "summary": "应整体回滚。", "content": {"mustPreserve": ["不应写入"]}},
            ]},
        }}
        run_id = self._start_run("验证跨操作事务回滚")
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        before = self.client.get(f"/api/v2/projects/{self.project_id}").json()
        response = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"],
            "selected_operation_ids": ["OP_METADATA_OK", "OP_METADATA_MISSING"],
            "expected_project_revision": run["base_project_revision"],
            "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"],
            "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
        })
        self.assertEqual(response.status_code, 404, response.text)
        self.assertEqual(self.client.get(f"/api/v2/projects/{self.project_id}").json(), before)
        self.assertEqual(self.client.get(f"/api/v2/agent/plans/{run['result']['plan_id']}").json()["status"], "awaiting_review")
        self.assertEqual(self.client.get(f"/api/v2/projects/{self.project_id}").json()["document"]["assets"][0].get("mustPreserve"), None)

    def test_provider_failure_keeps_user_message_attachment_and_run_error(self) -> None:
        self.adapter.error = server.ProviderError("Provider 暂时不可用", "network", 503)
        reference = self._upload("failure.mp4", b"local video reference", "video/mp4")
        run_id = self._start_run("模拟 Provider 失败但保留本地资料", [reference["id"]])
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "failed", run)
        self.assertEqual(run["error"]["kind"], "network")
        self.assertIn("Provider 暂时不可用", run["error"]["message"])
        self.assertEqual(self.client.get(reference["url"]).content, b"local video reference")
        messages = self.client.get(f"/api/v2/assistant/conversations/{run['conversation_id']}/messages").json()["messages"]
        self.assertEqual(len(messages), 1)
        self.assertEqual(messages[0]["role"], "user")
        self.assertEqual(messages[0]["attachments"][0]["id"], reference["id"])
        events = self.client.get(f"/api/v2/assistant/runs/{run_id}/events").json()["events"]
        self.assertTrue(any(item["event_type"] == "run_failed" for item in events))

    def test_prompt_claims_and_invalid_story_candidates_fail_server_validation(self) -> None:
        self.adapter.response = {"structured": {
            "reply": "不应接受伪造的 QA 状态。",
            "patch": {"workspace_operations": [{
                "id": "OP_FORBIDDEN_PROMPT", "workspace": "assets", "action": "create_prompt_candidate", "target_id": "AST_AGENT",
                "title": "伪造 QA", "summary": "应被后端拒绝。", "content": {"prompt": "角色正面肖像", "promptQaDecision": "Approved"},
            }]},
        }}
        prompt_run_id = self._start_run("拒绝伪造 Prompt QA")
        prompt_run = self._wait_for_terminal(prompt_run_id)
        self.assertEqual(prompt_run["status"], "failed", prompt_run)
        self.assertEqual(prompt_run["error"]["kind"], "validation")
        self.assertNotIn("plan_id", prompt_run["result"])

        self.adapter.response = {"structured": {
            "reply": "不应接受重复镜头 ID。",
            "patch": {"workspace_operations": [{
                "id": "OP_BAD_STORY", "workspace": "story", "action": "candidate_draft", "title": "重复镜头", "summary": "应被故事规范拒绝。",
                "content": {"script": "角色抬头。", "scenes": [{"id": "SC_1", "name": "夜"}], "shots": [
                    {"id": "SHOT_DUP", "scene": "SC_1", "duration": 2, "purpose": "验证", "size": "中景", "camera": "固定", "action": "抬头"},
                    {"id": "SHOT_DUP", "scene": "SC_1", "duration": 2, "purpose": "验证", "size": "中景", "camera": "固定", "action": "回望"},
                ]},
            }]},
        }}
        story_run_id = self._start_run("拒绝重复镜头候选")
        story_run = self._wait_for_terminal(story_run_id)
        self.assertEqual(story_run["status"], "failed", story_run)
        self.assertEqual(story_run["error"]["kind"], "validation")
        self.assertIn("镜头 ID 重复", story_run["error"]["message"])

    def test_assistant_resource_endpoints_and_terminal_run_guards(self) -> None:
        missing_paths = [
            "/api/v2/assistant/conversations/CONV_MISSING/messages",
            "/api/v2/assistant/attachments/ATT_MISSING",
            "/api/v2/assistant/runs/ARUN_MISSING",
            "/api/v2/assistant/runs/ARUN_MISSING/events",
            "/api/v2/assistant/conversations/CONV_MISSING/archive",
            "/api/v2/assistant/conversations/CONV_MISSING/restore",
            "/api/v2/assistant/conversations/CONV_MISSING/external-consent/reset",
        ]
        for path in missing_paths:
            response = self.client.get(path) if path.endswith("/messages") or path.endswith("/events") else self.client.post(path, json={}) if not path.endswith("/attachments/ATT_MISSING") and not path.endswith("/runs/ARUN_MISSING") else self.client.get(path)
            self.assertEqual(response.status_code, 404, f"{path}: {response.text}")
        run_id = self._start_run("验证终态运行的重复操作门禁")
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        duplicate_confirmation = self.client.post(f"/api/v2/assistant/runs/{run_id}/external-confirmation", json={"decision": "approve"})
        self.assertEqual(duplicate_confirmation.status_code, 409, duplicate_confirmation.text)
        duplicate_cancel = self.client.post(f"/api/v2/assistant/runs/{run_id}/cancel", json={})
        self.assertEqual(duplicate_cancel.status_code, 409, duplicate_cancel.text)

    def test_stream_rejects_project_path_mismatch_and_unknown_skill(self) -> None:
        mismatch = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": "PRJ_DIFFERENT", "message": "项目 ID 必须一致", "attachment_ids": [], "selected_node_ids": [],
            "skill_id": "video-script-storyboard", "context": {}, "cost_boundary": {}, "client_message_id": "mismatch-" + uuid.uuid4().hex,
        })
        self.assertEqual(mismatch.status_code, 409, mismatch.text)
        unknown_skill = self.client.post(f"/api/v2/projects/{self.project_id}/assistant/stream", json={
            "project_id": self.project_id, "message": "Skill 必须存在", "attachment_ids": [], "selected_node_ids": [],
            "skill_id": "skill-does-not-exist", "context": {}, "cost_boundary": {}, "client_message_id": "skill-" + uuid.uuid4().hex,
        })
        self.assertEqual(unknown_skill.status_code, 422, unknown_skill.text)

    def test_unsupported_extension_and_document_extraction_limit_are_explicit(self) -> None:
        unsupported = self.client.post(
            f"/api/v2/projects/{self.project_id}/assistant/attachments",
            files={"file": ("secret.zip", b"not supported", "application/zip")},
        )
        self.assertEqual(unsupported.status_code, 415, unsupported.text)
        source = self.temp_root / "bounded.txt"
        source.write_text("一二三四五六七八九十", encoding="utf-8")
        extracted = extract_document(source, source.name, "text/plain", maximum=5)
        self.assertEqual(extracted["status"], "succeeded")
        self.assertTrue(extracted["truncated"])
        self.assertIn("文档内容已按工作台抽取上限截断", extracted["text"])

    def test_workflow_candidate_applies_only_after_explicit_selection(self) -> None:
        self.adapter.response = {"structured": {
            "reply": "建议新增一个备注节点。",
            "patch": {
                "add_nodes": [{"id": "ASSISTANT_NOTE", "kind": "note", "label": "助手备注", "position": {"x": 42, "y": 42}}],
                "actions": ["node_orchestration"],
            },
        }}
        run_id = self._start_run("生成一个工作流图候选")
        run = self._wait_for_terminal(run_id)
        self.assertEqual(run["status"], "succeeded", run)
        operation = next(item for item in run["result"]["patch"]["workspace_operations"] if item["workspace"] == "workflow")
        self.assertEqual(operation["risk"], "review_required")
        graph_before = self.client.get(f"/api/v2/projects/{self.project_id}/graph").json()
        not_selected = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"], "selected_operation_ids": [],
            "expected_project_revision": run["base_project_revision"], "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"], "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
        })
        self.assertEqual(not_selected.status_code, 422, not_selected.text)
        applied = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"], "selected_operation_ids": [operation["id"]],
            "expected_project_revision": run["base_project_revision"], "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"], "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
        })
        self.assertEqual(applied.status_code, 200, applied.text)
        graph_after = self.client.get(f"/api/v2/projects/{self.project_id}/graph").json()
        self.assertEqual(graph_after["revision"], graph_before["revision"] + 1)
        self.assertIn("ASSISTANT_NOTE", {item["id"] for item in graph_after["graph"]["nodes"]})

    def test_apply_rejects_graph_and_timeline_revision_conflicts_without_write(self) -> None:
        run_id = self._start_run("验证图和时间线版本冲突")
        run = self._wait_for_terminal(run_id)
        operation_id = run["result"]["patch"]["workspace_operations"][0]["id"]
        contracts = self.client.get("/api/v2/contracts").json()
        graph_conflict = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"], "selected_operation_ids": [operation_id],
            "expected_project_revision": run["base_project_revision"], "expected_graph_revision": run["base_graph_revision"] + 1,
            "expected_timeline_revision": run["base_timeline_revision"], "expected_contract_bundle_hash": contracts["bundle_hash"],
        })
        self.assertEqual(graph_conflict.status_code, 409, graph_conflict.text)
        timeline_conflict = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"], "selected_operation_ids": [operation_id],
            "expected_project_revision": run["base_project_revision"], "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"] + 1, "expected_contract_bundle_hash": contracts["bundle_hash"],
        })
        self.assertEqual(timeline_conflict.status_code, 409, timeline_conflict.text)
        self.assertEqual(self.client.get(f"/api/v2/agent/plans/{run['result']['plan_id']}").json()["status"], "awaiting_review")

    def test_prompt_apply_preserves_active_asset_and_registered_history_fields(self) -> None:
        current = self.client.get(f"/api/v2/projects/{self.project_id}").json()
        protected = {
            "activeVersionId": "VER_ACTIVE", "activeArtifactId": "ART_ACTIVE", "qaDecision": "Approved",
            "regulatorRegistered": True, "generationStatus": "ready", "registrationStatus": "registered",
        }
        current["document"]["assets"][0].update(protected)
        saved = self.client.put(f"/api/v2/projects/{self.project_id}", json={
            "document": current["document"], "expected_revision": current["revision"],
        })
        self.assertEqual(saved.status_code, 200, saved.text)
        run_id = self._start_run("生成新的 Prompt 草稿但不能覆盖 active 资产")
        run = self._wait_for_terminal(run_id)
        operation_id = run["result"]["patch"]["workspace_operations"][0]["id"]
        applied = self.client.post(f"/api/v2/assistant/runs/{run_id}/apply", json={
            "plan_id": run["result"]["plan_id"], "selected_operation_ids": [operation_id],
            "expected_project_revision": run["base_project_revision"], "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"], "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
        })
        self.assertEqual(applied.status_code, 200, applied.text)
        asset = self.client.get(f"/api/v2/projects/{self.project_id}").json()["document"]["assets"][0]
        for key, value in protected.items():
            self.assertEqual(asset.get(key), value, key)
        self.assertEqual(asset["promptQaDecision"], "Pending")

    def test_docx_xlsx_and_pdf_are_extracted_locally(self) -> None:
        docx_path = self.temp_root / "reference.docx"
        docx_document = Document()
        docx_document.add_paragraph("角色身份锚点：左眉尾有浅色疤痕。")
        docx_document.save(docx_path)
        docx_result = extract_document(docx_path, docx_path.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert docx_result["status"] == "succeeded"
        assert "左眉尾有浅色疤痕" in docx_result["text"]

        xlsx_path = self.temp_root / "reference.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "资产表"
        worksheet.append(["资产 ID", "状态"])
        worksheet.append(["AST_AGENT", "保留历史版本"])
        workbook.save(xlsx_path)
        xlsx_result = extract_document(xlsx_path, xlsx_path.name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert xlsx_result["status"] == "succeeded"
        assert "AST_AGENT" in xlsx_result["text"]
        assert "保留历史版本" in xlsx_result["text"]

        pdf_path = self.temp_root / "reference.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 18 Tf 72 720 Td (PDF continuity note) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({
                NameObject("/F1"): writer._add_object(DictionaryObject({
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                })),
            }),
        })
        with pdf_path.open("wb") as handle:
            writer.write(handle)
        pdf_result = extract_document(pdf_path, pdf_path.name, "application/pdf")
        assert pdf_result["status"] == "succeeded"
        assert "PDF continuity note" in pdf_result["text"]

    def test_upload_rejects_mismatched_mime_without_leaving_a_file(self) -> None:
        rejected = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/attachments",
            files={"file": ("reference.png", b"not an image", "text/plain")},
        )
        assert rejected.status_code == 415, rejected.text
        assistant_dir = server.DATA_DIR / "projects" / self.project_id / "assistant" / "attachments"
        assert not list(assistant_dir.glob("*"))

    def test_document_extraction_and_external_confirmation_resume_same_run(self) -> None:
        source = self.temp_root / "notes.csv"
        source.write_text("镜头,动作\nS01,角色抬头\n", encoding="utf-8")
        extracted = extract_document(source, "notes.csv", "text/csv")
        assert extracted["status"] == "succeeded"
        assert "角色抬头" in extracted["text"]

        uploaded = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/attachments",
            files={"file": ("notes.csv", source.read_bytes(), "text/csv")},
        )
        assert uploaded.status_code == 200, uploaded.text
        attachment_id = uploaded.json()["attachment"]["id"]
        stream = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/stream",
            json={
                "project_id": self.project_id,
                "message": "根据资料生成角色 Prompt 候选",
                "attachment_ids": [attachment_id],
                "selected_node_ids": [],
                "skill_id": "video-character-design-director",
                "context": {"file_path": "/should/not/reach/provider"},
                "cost_boundary": {"confirmation_required": True},
                "client_message_id": "client-" + uuid.uuid4().hex,
            },
        )
        assert stream.status_code == 200, stream.text
        match = re.search(r'"run_id": "(ARUN_[^"]+)"', stream.text)
        assert match, stream.text
        run_id = match.group(1)
        assert self.adapter.calls == []
        run = self.client.get("/api/v2/assistant/runs/" + run_id).json()
        assert run["status"] == "awaiting_external_confirmation"
        assert run["awaiting_confirmation"]["attachments"][0]["delivery_mode"] == "extracted_text"

        approved = self.client.post(
            "/api/v2/assistant/runs/" + run_id + "/external-confirmation",
            json={"decision": "approve", "provider_profile_id": "opencode-default", "detail": {"approved_by": "test"}},
        )
        assert approved.status_code == 200, approved.text
        for _ in range(60):
            run = self.client.get("/api/v2/assistant/runs/" + run_id).json()
            if run["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert run["status"] == "succeeded", run
        assert len(self.adapter.calls) == 1
        assert "/should/not/reach/provider" not in self.adapter.calls[0][1]["input_text"]
        events = self.client.get("/api/v2/assistant/runs/" + run_id + "/events").json()["events"]
        assert [event["sequence"] for event in events] == list(range(1, len(events) + 1))
        assert any(event["event_type"] == "item_progress" for event in events)
        assert any(event["event_type"] == "run_completed" for event in events)

        plan_id = run["result"]["plan_id"]
        plan = self.client.get("/api/v2/agent/plans/" + plan_id).json()
        operation_ids = [item["id"] for item in plan["patch"]["workspace_operations"]]
        assert operation_ids == ["OP_AGENT_PROMPT"]
        applied = self.client.post(
            "/api/v2/assistant/runs/" + run_id + "/apply",
            json={
                "plan_id": plan_id,
                "selected_operation_ids": operation_ids,
                "expected_project_revision": run["base_project_revision"],
                "expected_graph_revision": run["base_graph_revision"],
                "expected_timeline_revision": run["base_timeline_revision"],
                "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
            },
        )
        assert applied.status_code == 200, applied.text
        assert applied.json()["applied_operation_ids"] == operation_ids
        project = self.client.get("/api/v2/projects/" + self.project_id).json()
        asset = project["document"]["assets"][0]
        assert asset["promptQaDecision"] == "Pending"
        assert asset["generationChoiceStatus"] == "user-confirmation-required"

    def test_image_is_sent_as_structured_vision_input_only_after_confirmation(self) -> None:
        self.adapter.capabilities.add("vision")
        uploaded = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/attachments",
            files={"file": ("reference.png", b"\x89PNG\r\n\x1a\nframeflow", "image/png")},
        )
        assert uploaded.status_code == 200, uploaded.text
        attachment_id = uploaded.json()["attachment"]["id"]
        stream = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/stream", json={
            "project_id": self.project_id,
            "message": "分析参考图并生成候选",
            "attachment_ids": [attachment_id],
            "selected_node_ids": [],
            "skill_id": "video-character-design-director",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "vision-" + uuid.uuid4().hex,
        })
        assert stream.status_code == 200, stream.text
        run_id = re.search(r'"run_id": "(ARUN_[^"]+)"', stream.text).group(1)
        assert self.adapter.calls == []
        awaiting = self.client.get("/api/v2/assistant/runs/" + run_id).json()
        assert awaiting["awaiting_confirmation"]["attachments"][0]["delivery_mode"] == "multimodal"
        approved = self.client.post("/api/v2/assistant/runs/" + run_id + "/external-confirmation", json={"decision": "approve", "provider_profile_id": "opencode-default", "detail": {"approved_by": "test"}})
        assert approved.status_code == 200, approved.text
        for _ in range(60):
            run = self.client.get("/api/v2/assistant/runs/" + run_id).json()
            if run["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert run["status"] == "succeeded", run
        assert [item[0] for item in self.adapter.calls] == ["vision", "orchestrator"]
        vision_input = self.adapter.calls[0][1]["input_content"]
        image_part = next(item for item in vision_input if item.get("type") == "input_image")
        assert image_part["image_url"].startswith("data:image/png;base64,")
        assert "/assistant/attachments/" not in image_part["image_url"]

    def test_cancel_keeps_an_unconfirmed_run_and_attachment_without_provider_call(self) -> None:
        uploaded = self.client.post(
            "/api/v2/projects/" + self.project_id + "/assistant/attachments",
            files={"file": ("cancel-note.txt", b"keep local", "text/plain")},
        )
        assert uploaded.status_code == 200, uploaded.text
        stream = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/stream", json={
            "project_id": self.project_id,
            "message": "分析附件",
            "attachment_ids": [uploaded.json()["attachment"]["id"]],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "cancel-" + uuid.uuid4().hex,
        })
        assert stream.status_code == 200, stream.text
        run_id = re.search(r'"run_id": "(ARUN_[^"]+)"', stream.text).group(1)
        canceled = self.client.post("/api/v2/assistant/runs/" + run_id + "/cancel", json={})
        assert canceled.status_code == 200, canceled.text
        assert canceled.json()["status"] == "canceled"
        assert self.adapter.calls == []
        events = self.client.get("/api/v2/assistant/runs/" + run_id + "/events").json()["events"]
        assert any(event["event_type"] == "run_interrupted" for event in events)
        assert self.client.get(uploaded.json()["attachment"]["url"]).content == b"keep local"

    def test_cancel_interrupts_a_running_provider_request_without_creating_a_plan(self) -> None:
        self.adapter.delay = 2.0
        uploaded = self._upload("slow-notes.txt", b"wait locally", "text/plain")
        run_id = self._start_run("启动一个可中止的慢运行", [uploaded["id"]])
        waiting = self.client.get(f"/api/v2/assistant/runs/{run_id}").json()
        self.assertEqual(waiting["status"], "awaiting_external_confirmation")
        approved = self.client.post(f"/api/v2/assistant/runs/{run_id}/external-confirmation", json={
            "decision": "approve",
            "provider_profile_id": "opencode-default",
            "detail": {"approved_by": "test"},
        })
        self.assertEqual(approved.status_code, 200, approved.text)
        canceled = self.client.post(f"/api/v2/assistant/runs/{run_id}/cancel", json={})
        self.assertEqual(canceled.status_code, 200, canceled.text)
        self.assertEqual(canceled.json()["status"], "canceled")
        time.sleep(0.1)
        final = self.client.get(f"/api/v2/assistant/runs/{run_id}").json()
        self.assertEqual(final["status"], "canceled")
        self.assertNotIn("plan_id", final["result"])
        self.assertEqual(self.client.get(uploaded["url"]).content, b"wait locally")
        events = self.client.get(f"/api/v2/assistant/runs/{run_id}/events").json()["events"]
        self.assertTrue(any(item["event_type"] == "run_interrupted" for item in events))

    def test_cross_workspace_apply_is_partial_and_preserves_unselected_graph(self) -> None:
        story = self.client.get("/api/v2/projects/" + self.project_id + "/story").json()
        timeline = self.client.get("/api/v2/projects/" + self.project_id + "/timeline").json()
        graph = self.client.get("/api/v2/projects/" + self.project_id + "/graph").json()
        story_document = story["story"]
        story_document["script"] = "夜色中的角色抬头，雨声暂时停下。"
        story_document["scenes"] = [{"id": "SC_AGENT", "name": "祠堂雨夜"}]
        story_document["shots"] = [{
            "id": "SHOT_AGENT",
            "scene": "SC_AGENT",
            "duration": 4,
            "purpose": "验证跨工作区候选",
            "size": "中景",
            "camera": "固定机位",
            "action": "角色抬头",
            "assetRequirements": [{"assetId": "AST_AGENT", "assetClass": "character", "role": "主角", "priority": "B", "required": True}],
        }]
        self.adapter.response = {"structured": {
            "reply": "已生成故事、资产元数据、时间线和工作流候选。",
            "patch": {
                "add_nodes": [{"id": "AGENT_NODE", "kind": "note", "label": "助手建议", "position": {"x": 40, "y": 40}}],
                "workspace_operations": [
                    {"id": "OP_STORY", "workspace": "story", "action": "candidate_draft", "title": "故事候选", "summary": "补充一条稳定镜头。", "content": story_document},
                    {"id": "OP_METADATA", "workspace": "assets", "action": "update_metadata", "target_id": "AST_AGENT", "title": "资产元数据候选", "summary": "增加身份保留说明。", "content": {"mustPreserve": ["角色脸部身份锚点"]}},
                    {"id": "OP_TIMELINE", "workspace": "timeline", "action": "candidate_draft", "title": "时间线候选", "summary": "保留当前轨道结构。", "content": {"document": timeline["document"]}},
                    {"id": "OP_AUDIO", "workspace": "audio", "action": "candidate_draft", "target_id": "AUD_AGENT", "title": "声音候选", "summary": "按镜头保留一条待确认朗读文本。", "content": {"audioDetails": {"sourceText": "我听见雨停了。", "textStatus": "candidate", "provider": "minimax"}}},
                ],
            },
            "actions": ["candidate_draft"],
            "next_skill": None,
            "requires_confirmation": False,
        }}
        stream = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/stream", json={
            "project_id": self.project_id,
            "message": "同步故事、资产元数据和时间线",
            "attachment_ids": [],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "multi-" + uuid.uuid4().hex,
        })
        assert stream.status_code == 200, stream.text
        run_id = re.search(r'"run_id": "(ARUN_[^"]+)"', stream.text).group(1)
        for _ in range(60):
            run = self.client.get("/api/v2/assistant/runs/" + run_id).json()
            if run["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert run["status"] == "succeeded", run
        plan_id = run["result"]["plan_id"]
        operations = self.client.get("/api/v2/agent/plans/" + plan_id).json()["patch"]["workspace_operations"]
        operation_ids = [item["id"] for item in operations]
        workflow_id = next(item["id"] for item in operations if item["workspace"] == "workflow")
        audio_operation = next(item for item in operations if item["id"] == "OP_AUDIO")
        assert all(field in audio_operation["content"]["audioDetails"] for field in self.client.get("/api/v2/contracts/audio").json()["required_fields"])
        selected = [item for item in operation_ids if item != workflow_id]
        applied = self.client.post("/api/v2/assistant/runs/" + run_id + "/apply", json={
            "plan_id": plan_id,
            "selected_operation_ids": selected,
            "expected_project_revision": run["base_project_revision"],
            "expected_graph_revision": run["base_graph_revision"],
            "expected_timeline_revision": run["base_timeline_revision"],
            "expected_contract_bundle_hash": self.client.get("/api/v2/contracts").json()["bundle_hash"],
        })
        assert applied.status_code == 200, applied.text
        assert applied.json()["applied_operation_ids"] == selected
        assert applied.json()["plan"]["status"] == "partially_applied"
        assert self.client.get("/api/v2/projects/" + self.project_id + "/graph").json()["revision"] == graph["revision"]
        refreshed_story = self.client.get("/api/v2/projects/" + self.project_id + "/story").json()["story"]
        assert refreshed_story["shots"][0]["id"] == "SHOT_AGENT"
        refreshed_project = self.client.get("/api/v2/projects/" + self.project_id).json()["document"]
        assert refreshed_project["assets"][0]["mustPreserve"] == ["角色脸部身份锚点"]

    def test_duplicate_client_message_is_idempotent(self) -> None:
        body = {
            "project_id": self.project_id,
            "message": "只检查项目状态",
            "attachment_ids": [],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "fixed-client-id",
        }
        first = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/stream", json=body)
        assert first.status_code == 200, first.text
        first_id = re.search(r'"run_id": "(ARUN_[^"]+)"', first.text).group(1)
        second = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/stream", json=body)
        assert second.status_code == 200, second.text
        second_id = re.search(r'"run_id": "(ARUN_[^"]+)"', second.text).group(1)
        assert first_id == second_id
        assert len(self.adapter.calls) == 1
        run = self.client.get(f"/api/v2/assistant/runs/{first_id}").json()
        messages = self.client.get(f"/api/v2/assistant/conversations/{run['conversation_id']}/messages").json()["messages"]
        assert len(messages) == 2
        assert messages[0]["client_message_id"] == "fixed-client-id"

    def test_contract_change_freezes_an_existing_plan(self) -> None:
        response = self.client.post("/api/v2/projects/" + self.project_id + "/assistant/stream", json={
            "project_id": self.project_id,
            "message": "生成一个候选计划",
            "attachment_ids": [],
            "selected_node_ids": [],
            "skill_id": "video-script-storyboard",
            "context": {},
            "cost_boundary": {},
            "client_message_id": "stale-" + uuid.uuid4().hex,
        })
        assert response.status_code == 200, response.text
        run_id = re.search(r'"run_id": "(ARUN_[^"]+)"', response.text).group(1)
        for _ in range(60):
            run = self.client.get("/api/v2/assistant/runs/" + run_id).json()
            if run["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert run["status"] == "succeeded", run
        changed_bundle = {**server.contract_bundle(), "bundle_hash": "f" * 64}
        with mock.patch.object(server, "contract_bundle", return_value=changed_bundle):
            blocked = self.client.post("/api/v2/assistant/runs/" + run_id + "/apply", json={
                "plan_id": run["result"]["plan_id"],
                "selected_operation_ids": ["OP_AGENT_PROMPT"],
                "expected_project_revision": run["base_project_revision"],
                "expected_graph_revision": run["base_graph_revision"],
                "expected_timeline_revision": run["base_timeline_revision"],
                "expected_contract_bundle_hash": run["contract_hash"],
            })
        assert blocked.status_code == 409, blocked.text
        stale_run = self.client.get("/api/v2/assistant/runs/" + run_id).json()
        assert stale_run["status"] == "stale_contract", stale_run
        assert stale_run["error"]["kind"] == "contract_stale", stale_run
        stale_plan = self.client.get("/api/v2/agent/plans/" + run["result"]["plan_id"]).json()
        assert stale_plan["status"] == "stale_contract", stale_plan
        assert stale_plan["decision"]["contract_stale"] is True, stale_plan
        assert stale_plan["decision"]["plan_contract_hash"] == run["contract_hash"], stale_plan
        assert stale_plan["decision"]["current_contract_hash"] == changed_bundle["bundle_hash"], stale_plan
